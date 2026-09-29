"""
Build a branded investor pitch deck (.pptx) from pitch Markdown.

The pitch is split on '## ' sections. Each section gets a layout chosen from
its content (table, stat callouts, numbered cards, or lead + body text), with
body text sized to fit and overflow continued on a follow-up slide.
"""
import math
import re
import threading
from datetime import date
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---------- Themes ----------
class Theme:
    FONT_KEYS = ('head_font', 'body_font')

    def __init__(self, key, name, **values):
        self.key, self.name = key, name
        for k, v in values.items():
            setattr(self, k, v if k in self.FONT_KEYS else RGBColor.from_string(v))


THEMES = {
    # Synoptic brand: dark cover/closing, white content slides, coral accents
    'coral': Theme(
        'coral', 'Synoptic Coral',
        accent='D44C3F', accent_strong='B73A2F', accent_soft='F19C92', tint='FEF7F6',
        ink='111827', body='374151', muted='6B7280', line='E5E7EB', panel='F9FAFB',
        bg='FFFFFF', on_accent='FFFFFF', cover_bg='1F232B', cover_text='FFFFFF', cover_muted='C9CDD4',
        head_font='Calibri', body_font='Calibri'),
    # Fresh and calm: forest-green cover, white content, mint accents, serif headings
    'green': Theme(
        'green', 'Modern Green',
        accent='2C7A4B', accent_strong='1F5C38', accent_soft='97D4A9', tint='EEF8F1',
        ink='13241A', body='3A4A40', muted='6B7A70', line='DDE7E0', panel='F5F9F6',
        bg='FFFFFF', on_accent='FFFFFF', cover_bg='173D28', cover_text='FFFFFF', cover_muted='B9D3C2',
        head_font='Cambria', body_font='Calibri'),
    # High contrast: dark slides throughout with vivid red accents
    'red': Theme(
        'red', 'Bold Red',
        accent='E5383B', accent_strong='FF6B6B', accent_soft='FF9A9C', tint='2A1718',
        ink='FFFFFF', body='D6D8DC', muted='9CA0A8', line='33363D', panel='1E2026',
        bg='131418', on_accent='FFFFFF', cover_bg='0B0B0D', cover_text='FFFFFF', cover_muted='A3A6AD',
        head_font='Arial', body_font='Arial'),
}
# Average character width (in ems) of each body font, for text fitting
CHAR_EM = {'Calibri': 0.47, 'Arial': 0.52}

DEFAULT_THEME = 'coral'
T = THEMES[DEFAULT_THEME]  # active theme while a deck is being built (see build_pitch_deck)

# ---------- Geometry (16:9 widescreen, inches) ----------
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6
CONTENT_TOP = 1.65
CONTENT_BOTTOM = 6.75
CONTENT_W = SLIDE_W - 2 * MARGIN
GAP = 0.35

MAX_BODY_PT, MIN_BODY_PT = 18, 12
READABLE_PT = 16  # below this, continue on another slide instead
NUMBER_RE = re.compile(r'(?<![\w.])([$€£]?\d[\d,.]*\s?(?:%|[xX]\b|[KkMmBbTt]\b|bn\b|million\b|billion\b)?\+?)')


# ============================================================
# Markdown parsing
# ============================================================
def parse_pitch(markdown):
    """Return (deck_title, [(section_title, [blocks])])."""
    text = (markdown or '').replace('\r', '')
    # Drop the trailing footer ("---" and anything after it)
    text = re.split(r'\n\s*---\s*\n', text)[0]
    title = None
    m = re.search(r'^#\s+(.+?)\s*$', text, re.M)
    if m:
        title = m.group(1).strip()
        text = text[:m.start()] + text[m.end():]

    sections, current = [], None
    for line in text.split('\n'):
        m = re.match(r'^##\s+(.+?)\s*$', line)
        if m and not line.startswith('###'):
            current = [m.group(1).strip(), []]
            sections.append(current)
        elif current is not None:
            current[1].append(line)
        elif line.strip():
            if not sections or sections[0][0] != 'Overview':
                sections.insert(0, ['Overview', []])
                current = sections[0]
            sections[0][1].append(line)
    return title, [(name, parse_blocks('\n'.join(lines))) for name, lines in sections if name]


def parse_blocks(body):
    """Split a section body into blocks: para, sub, list, table."""
    blocks, para = [], []

    def flush():
        if para:
            blocks.append({'type': 'para', 'text': ' '.join(para)})
            para.clear()

    lines = body.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            flush()
        elif re.match(r'^#{3,6}\s+', stripped):
            flush()
            blocks.append({'type': 'sub', 'text': re.sub(r'^#{3,6}\s+', '', stripped)})
        elif stripped.startswith('|'):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append({'type': 'table', 'rows': rows})
            continue
        elif re.match(r'^([-*+•▪●]|\d+[.)])\s+', stripped):
            flush()
            ordered = bool(re.match(r'^\d+[.)]\s+', stripped))
            items = []
            while i < len(lines) and re.match(r'^\s*([-*+•▪●]|\d+[.)])\s+', lines[i]):
                items.append(re.sub(r'^\s*([-*+•▪●]|\d+[.)])\s+', '', lines[i]).strip())
                i += 1
            blocks.append({'type': 'list', 'items': items, 'ordered': ordered})
            continue
        else:
            para.append(stripped)
        i += 1
    flush()
    return blocks


def plain(text):
    return re.sub(r'\*\*(.+?)\*\*|\*(.+?)\*|__(.+?)__', lambda m: next(g for g in m.groups() if g), text or '')


def inline_runs(text):
    """Split Markdown inline text into [(text, bold, italic)]."""
    runs, pos = [], 0
    for m in re.finditer(r'\*\*(.+?)\*\*|__(.+?)__|(?<![*\w])\*(?!\s)(.+?)\*(?![*\w])', text or ''):
        if m.start() > pos:
            runs.append((text[pos:m.start()], False, False))
        runs.append((m.group(1) or m.group(2) or m.group(3), bool(m.group(1) or m.group(2)), bool(m.group(3))))
        pos = m.end()
    if pos < len(text or ''):
        runs.append((text[pos:], False, False))
    return runs or [('', False, False)]


# ============================================================
# Text measurement (approximate, Calibri)
# ============================================================
def text_height(blocks, width_in, size_pt):
    """Estimate rendered height (inches) of text blocks at a font size."""
    char_w = size_pt * CHAR_EM.get(T.body_font, 0.52) / 72
    line_h = size_pt * 1.2 / 72
    per_line = max(10, int(width_in / char_w))
    height = 0.0
    for b in blocks:
        if b['type'] == 'para':
            height += math.ceil(len(plain(b['text'])) / per_line) * line_h + size_pt * 0.6 / 72
        elif b['type'] == 'sub':
            height += line_h * 1.15 + size_pt * 0.9 / 72
        elif b['type'] == 'list':
            item_chars = max(10, per_line - 3)
            for item in b['items']:
                height += math.ceil(len(plain(item)) / item_chars) * line_h + size_pt * 0.35 / 72
            height += size_pt * 0.4 / 72
    return height


def fit_size(blocks, width_in, height_in, floor=MIN_BODY_PT):
    for size in range(MAX_BODY_PT, floor - 1, -1):
        if text_height(blocks, width_in, size) <= height_in:
            return size
    return None


def split_blocks(blocks, width_in, height_in):
    """Split blocks into pages that fit at a readable size (splitting long lists)."""
    size = READABLE_PT
    pages, page = [], []
    expanded = []
    for b in blocks:
        if b['type'] == 'list' and text_height([b], width_in, size) > height_in * 0.6:
            for k in range(0, len(b['items']), 4):
                expanded.append({**b, 'items': b['items'][k:k + 4]})
        elif b['type'] == 'para' and text_height([b], width_in, size) > height_in * 0.6:
            chunk = ''
            for sentence in re.split(r'(?<=[.!?])\s+', b['text']):
                candidate = (chunk + ' ' + sentence).strip()
                if chunk and text_height([{'type': 'para', 'text': candidate}], width_in, size) > height_in * 0.5:
                    expanded.append({'type': 'para', 'text': chunk})
                    candidate = sentence
                chunk = candidate
            if chunk:
                expanded.append({'type': 'para', 'text': chunk})
        else:
            expanded.append(b)
    for b in expanded:
        if page and text_height(page + [b], width_in, size) > height_in:
            pages.append(page)
            page = []
        page.append(b)
    if page:
        pages.append(page)
    # A subheading shouldn't end a page
    for n in range(len(pages) - 1):
        if pages[n] and pages[n][-1]['type'] == 'sub' and len(pages[n]) > 1:
            pages[n + 1].insert(0, pages[n].pop())
    return pages


# ============================================================
# Drawing helpers
# ============================================================
def _rgb_hex(color):
    return '%02X%02X%02X' % (color[0], color[1], color[2])


def fill_background(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def add_box(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, radius=None):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def add_text(slide, x, y, w, h, text, size, color=None, bold=False, italic=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=None, markdown=False, line_spacing=None):
    color = color or T.body
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    for run_text, run_bold, run_italic in (inline_runs(text) if markdown else [(text, False, False)]):
        r = p.add_run()
        r.text = run_text
        _style(r, size, color, bold or run_bold, italic or run_italic, font)
    return box


def _style(run, size, color, bold=False, italic=False, font=None):
    run.font.name = font or T.body_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def _bullet(paragraph, char=None, autonum=False, color=None, indent_in=0.3):
    color = color or T.accent
    """Give a paragraph a real bullet (hanging indent)."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set('marL', str(Emu(Inches(indent_in))))
    pPr.set('indent', str(-Emu(Inches(indent_in))))
    for tag in ('a:buClr', 'a:buChar', 'a:buAutoNum', 'a:buNone', 'a:buFont'):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = pPr.makeelement(qn('a:buClr'), {})
    srgb = clr.makeelement(qn('a:srgbClr'), {'val': _rgb_hex(color)})
    clr.append(srgb)
    pPr.append(clr)
    if autonum:
        pPr.append(pPr.makeelement(qn('a:buAutoNum'), {'type': 'arabicPeriod'}))
    else:
        pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}))
        pPr.append(pPr.makeelement(qn('a:buChar'), {'char': char or '•'}))


def add_rich_text(slide, x, y, w, h, blocks, size):
    """Render para/sub/list blocks into one text box at the given size."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True

    def new_par():
        nonlocal first
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = 1.1
        return p

    for b in blocks:
        if b['type'] == 'para':
            p = new_par()
            p.space_after = Pt(size * 0.6)
            for t, bold, italic in inline_runs(b['text']):
                r = p.add_run(); r.text = t
                _style(r, size, T.ink if bold else T.body, bold, italic)
        elif b['type'] == 'sub':
            p = new_par()
            p.space_before = Pt(size * 0.4)
            p.space_after = Pt(size * 0.3)
            r = p.add_run(); r.text = plain(b['text'])
            _style(r, size + 2, T.accent_strong, True, font=T.head_font)
        elif b['type'] == 'list':
            for n, item in enumerate(b['items']):
                p = new_par()
                p.space_after = Pt(size * 0.35)
                _bullet(p, autonum=b['ordered'])
                for t, bold, italic in inline_runs(item):
                    r = p.add_run(); r.text = t
                    _style(r, size, T.ink if bold else T.body, bold, italic)
            p.space_after = Pt(size * 0.75)
    return box


def _cell_borders(cell, color, width_pt=0.75):
    """Give a table cell thin borders in the theme's line colour."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
        ln = tcPr.makeelement(qn(tag), {'w': str(int(width_pt * 12700)), 'cap': 'flat', 'cmpd': 'sng', 'algn': 'ctr'})
        fill = ln.makeelement(qn('a:solidFill'), {})
        fill.append(fill.makeelement(qn('a:srgbClr'), {'val': _rgb_hex(color)}))
        ln.append(fill)
        ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'solid'}))
        # Borders must precede the cell fill in tcPr
        tcPr.insert(['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'].index(tag), ln)


def _shrink_on_overflow(box):
    """Ask PowerPoint to shrink text that still doesn't fit its box."""
    bodyPr = box.text_frame._txBody.bodyPr
    for tag in ('a:noAutofit', 'a:spAutoFit', 'a:normAutofit'):
        for el in bodyPr.findall(qn(tag)):
            bodyPr.remove(el)
    bodyPr.append(bodyPr.makeelement(qn('a:normAutofit'), {}))


def number_badge(slide, x, y, d, label, fill=None, color=None, size=None):
    fill, color = fill or T.accent, color or T.on_accent
    circle = add_box(slide, x, y, d, d, fill, MSO_SHAPE.OVAL)
    tf = circle.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    _style(r, size or d * 30, color, True)
    return circle


# ============================================================
# Slide builders
# ============================================================
class DeckBuilder:
    def __init__(self, deck_title, subtitle='Investor Pitch Deck'):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.blank = self.prs.slide_layouts[6]
        self.title = deck_title
        self.subtitle = subtitle
        self.page = 0
        self.one_slide_per_section = False  # set for fixed-length decks

    def _new_slide(self, notes=None):
        slide = self.prs.slides.add_slide(self.blank)
        self.page += 1
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    def _chrome(self, slide, number, title):
        """Section badge + title + footer on a light content slide."""
        fill_background(slide, T.bg)
        number_badge(slide, MARGIN, 0.62, 0.62, f'{number:02d}', size=15)
        add_text(slide, MARGIN + 0.85, 0.5, CONTENT_W - 0.85, 0.85, title, 30, T.ink, bold=True,
                 anchor=MSO_ANCHOR.MIDDLE, font=T.head_font)
        add_text(slide, MARGIN, 7.0, 8, 0.3, self.title, 10, T.muted)
        add_text(slide, SLIDE_W - MARGIN - 1, 7.0, 1, 0.3, str(self.page), 10, T.muted, align=PP_ALIGN.RIGHT)

    # ---- cover / agenda / closing ----
    def cover(self):
        slide = self._new_slide()
        fill_background(slide, T.cover_bg)
        # Motif: overlapping coral circles bleeding off the right edge
        add_box(slide, 8.9, -1.2, 6.2, 6.2, T.accent, MSO_SHAPE.OVAL)
        add_box(slide, 10.6, 3.9, 4.2, 4.2, T.accent_strong, MSO_SHAPE.OVAL)
        add_box(slide, 8.2, 4.7, 1.3, 1.3, T.accent_soft, MSO_SHAPE.OVAL)
        add_text(slide, MARGIN + 0.1, 1.2, 7.6, 0.4, self.subtitle.upper(), 14, T.accent_soft, bold=True)
        size = 54 if len(self.title) <= 22 else 44 if len(self.title) <= 36 else 36
        add_text(slide, MARGIN + 0.1, 1.75, 7.6, 3.2, self.title, size, T.cover_text, bold=True, font=T.head_font,
                 anchor=MSO_ANCHOR.TOP, line_spacing=0.95)
        add_text(slide, MARGIN + 0.1, 6.35, 7, 0.4, date.today().strftime('%B %Y'), 14, T.cover_muted)

    def agenda(self, section_names):
        slide = self._new_slide()
        fill_background(slide, T.bg)
        add_text(slide, MARGIN, 0.55, 5, 0.9, 'Agenda', 36, T.ink, bold=True, font=T.head_font)
        add_text(slide, MARGIN, 1.4, 4.4, 1.2, f'What we\'ll cover in the next {len(section_names)} sections.', 16, T.muted)
        # Two columns of numbered rows on the right
        names = section_names[:20]
        per_col = math.ceil(len(names) / 2) if len(names) > 5 else len(names)
        col_w = 3.7 if len(names) > 5 else 6.6
        row_h = min(0.95, 5.6 / max(per_col, 1))
        badge = min(0.55, row_h - 0.1)
        name_size = 18 if row_h >= 0.8 else 15 if row_h >= 0.6 else 13
        for i, name in enumerate(names):
            col, row = divmod(i, per_col)
            x = 5.35 + col * (col_w + 0.35)
            y = 0.95 + row * row_h
            number_badge(slide, x, y + (row_h - badge) / 2, badge, f'{i + 1:02d}', fill=T.tint, color=T.accent_strong,
                         size=14 if badge >= 0.5 else 11)
            add_text(slide, x + badge + 0.2, y, col_w - badge - 0.25, row_h, name, name_size, T.ink, bold=True,
                     anchor=MSO_ANCHOR.MIDDLE)
        add_text(slide, MARGIN, 7.0, 8, 0.3, self.title, 10, T.muted)
        add_text(slide, SLIDE_W - MARGIN - 1, 7.0, 1, 0.3, str(self.page), 10, T.muted, align=PP_ALIGN.RIGHT)

    def closing(self):
        slide = self._new_slide()
        fill_background(slide, T.cover_bg)
        add_box(slide, -1.6, 3.6, 5.2, 5.2, T.accent, MSO_SHAPE.OVAL)
        add_box(slide, 11.4, -1.0, 3.0, 3.0, T.accent_strong, MSO_SHAPE.OVAL)
        add_text(slide, 3.0, 2.35, 7.33, 1.2, 'Thank you', 54, T.cover_text, bold=True, font=T.head_font, align=PP_ALIGN.CENTER)
        name_size = 22 if len(self.title) <= 40 else 18
        add_text(slide, 2.5, 3.6, 8.33, 0.9, self.title, name_size, T.accent_soft, align=PP_ALIGN.CENTER)
        add_text(slide, 3.0, 4.65, 7.33, 0.5, 'Questions & discussion', 16, T.cover_muted,
                 align=PP_ALIGN.CENTER)

    # ---- section content ----
    def section(self, number, name, blocks):
        notes = section_notes(name, blocks)
        layout = choose_layout(blocks)
        if layout == 'table':
            return self._table_slide(number, name, blocks, notes)
        if layout == 'stats':
            return self._stats_slide(number, name, blocks, notes)
        if layout == 'cards':
            return self._cards_slide(number, name, blocks, notes)
        return self._text_slides(number, name, blocks, notes)

    def _text_slides(self, number, name, blocks, notes):
        """Lead statement on the left, body on the right; continues if too long."""
        lead, rest = split_lead(blocks)
        height = CONTENT_BOTTOM - CONTENT_TOP
        if self.one_slide_per_section and lead and not fit_size(rest, CONTENT_W - 4.55, height):
            lead, rest = None, blocks  # give the body the full width rather than overflow
        body_x = MARGIN + 4.55 if lead else MARGIN
        body_w = SLIDE_W - MARGIN - body_x
        if self.one_slide_per_section:
            # Fixed-length deck: everything on one slide, shrinking text as needed
            size = (fit_size(rest, body_w, height) or MIN_BODY_PT) if rest else MAX_BODY_PT
            pages = [rest]
        else:
            size = fit_size(rest, body_w, height, READABLE_PT) if rest else MAX_BODY_PT
            pages = [rest] if size else split_blocks(rest, body_w, height)
        if not self.one_slide_per_section and len(pages) == 2 and text_height(pages[1], body_w, READABLE_PT) < height * 0.3:
            # Avoid a near-empty continuation slide: allow one point smaller instead
            tighter = fit_size(rest, body_w, height, READABLE_PT - 1)
            if tighter:
                size, pages = tighter, [rest]
        if len(pages) > 1 and lead:
            # Continuation slides have no lead box, so re-flow them at full width
            pages = pages[:1] + split_blocks([b for pg in pages[1:] for b in pg], CONTENT_W, height)
        for n, page in enumerate(pages):
            slide = self._new_slide(notes)
            self._chrome(slide, number, name if n == 0 else f'{name} (cont.)')
            if lead and n == 0:
                add_box(slide, MARGIN, CONTENT_TOP, 4.15, height, T.tint, MSO_SHAPE.ROUNDED_RECTANGLE, 0.06)
                lead_size = 26 if len(lead) < 110 else 22 if len(lead) < 180 else 19
                add_text(slide, MARGIN + 0.35, CONTENT_TOP + 0.35, 3.45, height - 0.7, lead, lead_size, T.ink,
                         bold=True, font=T.head_font, anchor=MSO_ANCHOR.MIDDLE, markdown=True, line_spacing=1.1)
            elif lead:
                body_x, body_w = MARGIN, CONTENT_W
            if page:
                page_size = size or READABLE_PT  # one size across continuation slides
                box = add_rich_text(slide, body_x, CONTENT_TOP + 0.05, body_w, height, page, page_size)
                if self.one_slide_per_section and not fit_size(page, body_w, height):
                    _shrink_on_overflow(box)

    def _cards_slide(self, number, name, blocks, notes):
        intro = [b for b in blocks if b['type'] in ('para', 'sub')]
        items = next(b for b in blocks if b['type'] == 'list')['items']
        slide = self._new_slide(notes)
        self._chrome(slide, number, name)
        top = CONTENT_TOP
        if intro:
            size = fit_size(intro, CONTENT_W, 1.3) or 14
            h = min(1.3, text_height(intro, CONTENT_W, size) + 0.05)
            add_rich_text(slide, MARGIN, top, CONTENT_W, h, intro, size)
            top += h + 0.3
        cols = len(items) if len(items) <= 3 else 2 if len(items) == 4 else 3
        rows = math.ceil(len(items) / cols)
        card_w = (CONTENT_W - GAP * (cols - 1)) / cols
        max_card_h = (CONTENT_BOTTOM - top - GAP * (rows - 1)) / rows
        text_w = card_w - 0.7
        size = min(fit_size([{'type': 'para', 'text': it}], text_w, max_card_h - 1.4) or MIN_BODY_PT for it in items)
        size = min(size, 20)
        needed = max(text_height([{'type': 'para', 'text': it}], text_w, size) for it in items)
        card_h = min(max_card_h, max(2.0, needed + 1.5))
        text_h = card_h - 1.15
        block_h = rows * card_h + GAP * (rows - 1)
        top += max(0, (CONTENT_BOTTOM - top - block_h) / 2 - 0.25)  # centre the card group
        for i, item in enumerate(items):
            r, c = divmod(i, cols)
            x = MARGIN + c * (card_w + GAP)
            y = top + r * (card_h + GAP)
            add_box(slide, x, y, card_w, card_h, T.panel, MSO_SHAPE.ROUNDED_RECTANGLE, 0.08)
            number_badge(slide, x + 0.35, y + 0.35, 0.5, str(i + 1), size=14)
            add_text(slide, x + 0.35, y + 1.05, text_w, text_h, item, size, T.body, markdown=True, line_spacing=1.1)

    def _stats_slide(self, number, name, blocks, notes):
        intro = [b for b in blocks if b['type'] in ('para', 'sub')]
        items = next(b for b in blocks if b['type'] == 'list')['items']
        slide = self._new_slide(notes)
        self._chrome(slide, number, name)
        top = CONTENT_TOP
        if intro:
            size = fit_size(intro, CONTENT_W, 1.4) or 14
            h = min(1.4, text_height(intro, CONTENT_W, size) + 0.05)
            add_rich_text(slide, MARGIN, top, CONTENT_W, h, intro, size)
            top += h + 0.35
        cols = len(items)
        card_w = (CONTENT_W - GAP * (cols - 1)) / cols
        labels = [extract_stat(it)[1] for it in items]
        label_size = min(16, min(fit_size([{'type': 'para', 'text': l}], card_w - 0.7, CONTENT_BOTTOM - top - 1.9)
                                 or MIN_BODY_PT for l in labels))
        needed = max(text_height([{'type': 'para', 'text': l}], card_w - 0.7, label_size) for l in labels)
        card_h = min(CONTENT_BOTTOM - top, max(2.8, needed + 2.0))
        top += max(0, (CONTENT_BOTTOM - top - card_h) / 2 - 0.25)  # centre the stat cards
        for i, item in enumerate(items):
            figure, label = extract_stat(item)
            x = MARGIN + i * (card_w + GAP)
            add_box(slide, x, top, card_w, card_h, T.tint, MSO_SHAPE.ROUNDED_RECTANGLE, 0.06)
            fig_size = 44 if len(figure) <= 6 else 36 if len(figure) <= 9 else 28
            add_text(slide, x + 0.35, top + 0.35, card_w - 0.7, 1.0, figure, fig_size, T.accent_strong, bold=True, font=T.head_font,
                     anchor=MSO_ANCHOR.BOTTOM)
            label_h = card_h - 1.65
            add_text(slide, x + 0.35, top + 1.5, card_w - 0.7, label_h, label, label_size, T.body,
                     markdown=True, line_spacing=1.1)

    def _table_slide(self, number, name, blocks, notes):
        table = next(b for b in blocks if b['type'] == 'table')
        others = [b for b in blocks if b is not table and b['type'] != 'table']
        slide = self._new_slide(notes)
        self._chrome(slide, number, name)
        height = CONTENT_BOTTOM - CONTENT_TOP
        if others:
            side_w = 4.3
            size = fit_size(others, side_w, height) or MIN_BODY_PT
            add_rich_text(slide, MARGIN, CONTENT_TOP + 0.05, side_w, height, others, size)
            tx, tw = MARGIN + side_w + 0.5, CONTENT_W - side_w - 0.5
        else:
            tx, tw = MARGIN, CONTENT_W
        rows = table['rows'][:12]
        ncols = max(len(r) for r in rows)
        row_h = min(0.6, height / len(rows))
        shape = slide.shapes.add_table(len(rows), ncols, Inches(tx), Inches(CONTENT_TOP),
                                       Inches(tw), Inches(row_h * len(rows)))
        tbl = shape.table
        # Plain styling: no banding from the default table style
        tblPr = tbl._tbl.tblPr
        tblPr.set('bandRow', '0')
        tblPr.set('firstRow', '1')
        size = 16 if len(rows) <= 6 else 13
        for r, row in enumerate(rows):
            tbl.rows[r].height = Inches(row_h)
            for c in range(ncols):
                cell = tbl.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = T.accent if r == 0 else (T.bg if r % 2 else T.panel)
                _cell_borders(cell, T.line)
                cell.margin_left = cell.margin_right = Inches(0.15)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                text = row[c] if c < len(row) else ''
                for t, bold, italic in inline_runs(text):
                    run = p.add_run(); run.text = t
                    _style(run, size, T.on_accent if r == 0 else (T.ink if c == 0 or bold else T.body),
                           r == 0 or c == 0 or bold, italic)


# ============================================================
# Layout selection
# ============================================================
def choose_layout(blocks):
    types = [b['type'] for b in blocks]
    lists = [b for b in blocks if b['type'] == 'list']
    paras = [b for b in blocks if b['type'] == 'para']
    intro_chars = sum(len(plain(b['text'])) for b in paras)
    if 'table' in types and types.count('table') == 1:
        return 'table'
    if 'sub' in types or len(lists) != 1:
        return 'text'
    items = lists[0]['items']
    if intro_chars > 320:
        return 'text'
    if 2 <= len(items) <= 4 and all(extract_stat(i)[0] for i in items) and all(len(plain(i)) <= 110 for i in items):
        return 'stats'
    if 2 <= len(items) <= 6 and all(len(plain(i)) <= 170 for i in items):
        return 'cards'
    return 'text'


def extract_stat(item):
    """Return (figure, label) when a list item leads with / contains a key number."""
    text = plain(item)
    # "Label: $1.2B" or "TAM - $1.2B"
    m = re.match(r'^\s*([^:–—-]{1,40}?)\s*[:–—-]\s*(.+)$', text)
    if m:
        num = NUMBER_RE.search(m.group(2))
        if num and num.start() == 0:
            figure = num.group(1).strip()
            rest = m.group(2)[num.end():].strip(' ,.;')
            return figure, (m.group(1) + (' — ' + rest if rest else '')).strip()
    num = NUMBER_RE.search(text)
    if num and num.start() <= 3 and re.search(r'[%$€£KkMmBbxX]|\d{2,}', num.group(1)):
        figure = num.group(1).strip()
        label = (text[:num.start()] + text[num.end():]).strip(' ,.;:-–—')
        label = re.sub(r'^(of|in|to)\s+', '', label)
        return figure, label[:1].upper() + label[1:]
    return '', text


def split_lead(blocks):
    """Use the first sentence of the first paragraph as a lead statement."""
    if not blocks or blocks[0]['type'] != 'para':
        return None, blocks
    text = blocks[0]['text']
    m = re.match(r'^(.{25,150}?[.!?])(\s+|$)(.*)$', text, re.S)
    if not m:
        if len(text) <= 150 and blocks[1:]:
            return text, blocks[1:]
        return None, blocks
    lead, remainder = m.group(1), m.group(3).strip()
    rest = ([{'type': 'para', 'text': remainder}] if remainder else []) + blocks[1:]
    if not rest:
        return None, blocks  # a single short sentence reads better as body text
    return lead, rest


def section_notes(name, blocks):
    lines = [name]
    for b in blocks:
        if b['type'] in ('para', 'sub'):
            lines.append(plain(b['text']))
        elif b['type'] == 'list':
            lines.extend('- ' + plain(i) for i in b['items'])
        elif b['type'] == 'table':
            lines.extend(' | '.join(plain(c) for c in r) for r in b['rows'])
    return '\n'.join(lines)


# ============================================================
# Public API
# ============================================================
_theme_lock = threading.Lock()


def build_pitch_deck(project_title, markdown, theme=DEFAULT_THEME, slide_count=None):
    """
    Return a BytesIO containing the finished .pptx in the given theme.
    With slide_count, the deck follows that length plan exactly (see slide_plan.py).
    """
    global T
    with _theme_lock:  # T is module state; keep concurrent exports from mixing themes
        T = THEMES.get(theme, THEMES[DEFAULT_THEME])
        try:
            return _build(project_title, markdown, slide_count)
        finally:
            T = THEMES[DEFAULT_THEME]


def fold_sections(sections, limit):
    """Merge trailing extra sections into the one before them (as subsections)."""
    sections = [(n, list(b)) for n, b in sections]
    while len(sections) > limit > 0:
        name, blocks = sections.pop()
        sections[-1][1].extend([{'type': 'sub', 'text': name}] + blocks)
    return sections


def _build(project_title, markdown, slide_count=None):
    deck_title, sections = parse_pitch(markdown)
    title = (project_title or deck_title or 'Pitch Deck').strip()
    title = re.sub(r'\s+pitch deck$', '', title, flags=re.I) or title
    builder = DeckBuilder(title)
    sections = [(n, b) for n, b in sections if b] or [('Overview', [{'type': 'para', 'text': 'No pitch content yet.'}])]

    if slide_count:
        from slide_plan import PLANS
        builder.one_slide_per_section = True
        plan = PLANS.get(int(slide_count))
        if plan:
            sections = fold_sections(sections, len(plan['sections']))
            show_agenda, show_closing = plan['agenda'], plan['closing']
        else:
            # Fixed-length pitch from an older plan: keep its sections, one slide each
            show_agenda, show_closing = len(sections) >= 8, True
    else:
        show_agenda, show_closing = len(sections) >= 3, True

    builder.cover()
    if show_agenda:
        builder.agenda([n for n, _ in sections])
    for i, (name, blocks) in enumerate(sections, start=1):
        builder.section(i, name, blocks)
    if show_closing:
        builder.closing()
    buffer = BytesIO()
    builder.prs.save(buffer)
    buffer.seek(0)
    return buffer
