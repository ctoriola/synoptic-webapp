import os
import re
import json
from urllib.parse import urlparse
from datetime import datetime
from io import BytesIO

import requests
from flask import Blueprint, request, jsonify, send_file, redirect, url_for, session, current_app
from flask_login import login_required, current_user
# Replaced Gemini with Hugging Face for free inference
from huggingface_client import generate_pitch_json, refine_pitch_content, generate_basic_template, deep_pitch_generation
# PDF generation temporarily disabled for Vercel compatibility
# from reportlab.lib.pagesizes import letter, A4
# from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
# from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
# from reportlab.lib.units import inch
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

from firebase_models import Project, User

api_bp = Blueprint('api', __name__)

def search_web_content(query, max_results=3):
    """Search the web for relevant content using a search API"""
    try:
        # Use DuckDuckGo Instant Answer API as a fallback search
        search_url = "https://api.duckduckgo.com/"
        params = {
            'q': query,
            'format': 'json',
            'no_redirect': '1',
            'no_html': '1',
            'skip_disambig': '1'
        }
        
        response = requests.get(search_url, params=params, timeout=10)
        if response.status_code == 200:
            data = response.json()
            results = []
            
            # Extract abstract and related topics
            if data.get('Abstract'):
                results.append({
                    'title': data.get('AbstractText', 'Market Research'),
                    'snippet': data.get('Abstract'),
                    'url': data.get('AbstractURL', '')
                })
            
            # Extract related topics
            for topic in data.get('RelatedTopics', [])[:max_results-1]:
                if isinstance(topic, dict) and topic.get('Text'):
                    results.append({
                        'title': topic.get('Text', '')[:100],
                        'snippet': topic.get('Text', ''),
                        'url': topic.get('FirstURL', '')
                    })
            
            return results[:max_results]
            
    except Exception as e:
        print(f"DEBUG: Web search failed: {str(e)}")
        
    # Fallback: return mock data based on query keywords
    return generate_mock_market_data(query)

def generate_mock_market_data(query):
    """Generate realistic mock market data when web search fails"""
    mock_data = []
    
    if "market size" in query.lower():
        mock_data.append({
            'title': 'Market Size Analysis',
            'snippet': f'The global market for {query.split()[0]} technology is projected to reach significant growth, with increasing adoption across various industries and expanding user base.',
            'url': ''
        })
    elif "trends" in query.lower():
        mock_data.append({
            'title': 'Industry Trends',
            'snippet': f'Current trends in {query.split()[0]} technology show accelerating innovation, increased investment, and growing market demand driven by digital transformation initiatives.',
            'url': ''
        })
    elif "competitors" in query.lower():
        mock_data.append({
            'title': 'Competitive Landscape',
            'snippet': f'The competitive landscape for {query.split()[0]} includes both established players and emerging startups, with differentiation opportunities in user experience and technical innovation.',
            'url': ''
        })
    else:
        mock_data.append({
            'title': 'Market Research',
            'snippet': f'Market analysis indicates growing opportunities in the {query.split()[0]} sector with favorable conditions for new entrants and innovative solutions.',
            'url': ''
        })
    
    return mock_data

# Configuration
# Removed GOOGLE_API_KEY - now using Hugging Face
GITHUB_TOKEN = os.getenv("GITHUB_TOKEN")
HF_API_TOKEN = os.getenv("HF_API_TOKEN")

# Constants from original app
essential_readme_paths = [
    "README.md", "README.MD", "Readme.md", "README",
    "docs/README.md", "docs/README"
]

common_branches = ["main", "master", "develop"]

def parse_github_repo(url: str):
    """Return (owner, repo) if URL looks like a GitHub repo; else (None, None)."""
    if not url:
        return None, None
    try:
        u = urlparse(url)
    except Exception:
        return None, None

    host = (u.netloc or '').lower()
    if host not in {"github.com", "www.github.com"}:
        return None, None

    parts = [p for p in (u.path or '').split('/') if p]
    if len(parts) < 2:
        return None, None

    owner, repo = parts[0], parts[1]
    if repo.endswith('.git'):
        repo = repo[:-4]
    return owner, repo

def try_fetch_readme_raw(owner: str, repo: str):
    """Try common branches and README paths from raw.githubusercontent.com."""
    for br in common_branches:
        for p in essential_readme_paths:
            raw = f"https://raw.githubusercontent.com/{owner}/{repo}/{br}/{p}"
            try:
                r = requests.get(raw, timeout=12)
                if r.status_code == 200 and r.text.strip():
                    return r.text, raw
            except requests.RequestException:
                continue
    return None, None

def try_fetch_readme_api(owner: str, repo: str, user_token: str = None):
    """Fallback to GitHub API to get the default README if possible."""
    api = f"https://api.github.com/repos/{owner}/{repo}/readme"
    headers = {
        "Accept": "application/vnd.github+json",
        "User-Agent": "PitchPerfectAI-saas",
    }
    # Use user's GitHub token first (for private repos), fallback to global token
    token = user_token or GITHUB_TOKEN
    if token:
        headers["Authorization"] = f"Bearer {token}"
    try:
        r = requests.get(api, headers=headers, timeout=12)
        if r.status_code == 200:
            j = r.json()
            download_url = j.get("download_url")
            if download_url:
                rr = requests.get(download_url, timeout=12)
                if rr.status_code == 200 and rr.text.strip():
                    return rr.text, download_url
    except requests.RequestException:
        pass
    return None, None

def analyze_repository_structure(owner: str, repo: str, user_token: str = None) -> dict:
    """Analyze repository structure and code to extract project insights without README."""
    headers = {
        "Accept": "application/vnd.github+json",
        "User-Agent": "PitchPerfectAI-saas",
    }
    if user_token:
        headers["Authorization"] = f"Bearer {user_token}"
    
    analysis = {
        "project_type": "Unknown",
        "tech_stack": [],
        "main_language": "Unknown",
        "frameworks": [],
        "features": [],
        "structure_insights": [],
        "dependencies": {},
        "file_count": 0,
        "directory_structure": []
    }
    
    try:
        # Get repository contents
        api_url = f"https://api.github.com/repos/{owner}/{repo}/contents"
        response = requests.get(api_url, headers=headers, timeout=10)
        
        if response.status_code == 200:
            contents = response.json()
            analysis["file_count"] = len(contents)
            
            # Analyze root directory structure
            for item in contents:
                if item["type"] == "dir":
                    analysis["directory_structure"].append(item["name"])
                elif item["type"] == "file":
                    filename = item["name"].lower()
                    
                    # Detect project type and tech stack from files
                    if filename == "package.json":
                        analysis["project_type"] = "Node.js/JavaScript"
                        analysis["tech_stack"].append("Node.js")
                        # Try to get package.json content for dependencies
                        try:
                            pkg_response = requests.get(item["download_url"], timeout=5)
                            if pkg_response.status_code == 200:
                                pkg_data = pkg_response.json()
                                if "dependencies" in pkg_data:
                                    analysis["dependencies"]["runtime"] = list(pkg_data["dependencies"].keys())[:10]
                                if "devDependencies" in pkg_data:
                                    analysis["dependencies"]["dev"] = list(pkg_data["devDependencies"].keys())[:10]
                        except:
                            pass
                    
                    elif filename == "requirements.txt":
                        analysis["project_type"] = "Python"
                        analysis["tech_stack"].append("Python")
                        # Try to get requirements content
                        try:
                            req_response = requests.get(item["download_url"], timeout=5)
                            if req_response.status_code == 200:
                                deps = [line.split("==")[0].split(">=")[0].split("<=")[0].strip() 
                                       for line in req_response.text.split("\n") if line.strip()]
                                analysis["dependencies"]["python"] = deps[:15]
                        except:
                            pass
                    
                    elif filename == "pom.xml":
                        analysis["project_type"] = "Java/Maven"
                        analysis["tech_stack"].extend(["Java", "Maven"])
                    
                    elif filename == "build.gradle":
                        analysis["project_type"] = "Java/Gradle"
                        analysis["tech_stack"].extend(["Java", "Gradle"])
                    
                    elif filename == "cargo.toml":
                        analysis["project_type"] = "Rust"
                        analysis["tech_stack"].append("Rust")
                    
                    elif filename == "go.mod":
                        analysis["project_type"] = "Go"
                        analysis["tech_stack"].append("Go")
                    
                    elif filename == "composer.json":
                        analysis["project_type"] = "PHP"
                        analysis["tech_stack"].append("PHP")
                    
                    elif filename in ["dockerfile", "docker-compose.yml", "docker-compose.yaml"]:
                        analysis["tech_stack"].append("Docker")
                    
                    elif filename in [".github", ".gitlab-ci.yml", "jenkinsfile"]:
                        analysis["features"].append("CI/CD Pipeline")
                    
                    elif filename in ["vercel.json", "netlify.toml"]:
                        analysis["features"].append("Cloud Deployment")
        
        # Get language statistics
        lang_url = f"https://api.github.com/repos/{owner}/{repo}/languages"
        lang_response = requests.get(lang_url, headers=headers, timeout=10)
        if lang_response.status_code == 200:
            languages = lang_response.json()
            if languages:
                analysis["main_language"] = max(languages.keys(), key=lambda k: languages[k])
                analysis["tech_stack"].extend(list(languages.keys())[:5])
        
        # Detect frameworks based on directory structure and dependencies
        dirs = analysis["directory_structure"]
        deps = analysis["dependencies"]
        
        # Web frameworks detection
        if "react" in str(deps).lower() or "src" in dirs:
            analysis["frameworks"].append("React")
        if "vue" in str(deps).lower():
            analysis["frameworks"].append("Vue.js")
        if "angular" in str(deps).lower():
            analysis["frameworks"].append("Angular")
        if "express" in str(deps).lower():
            analysis["frameworks"].append("Express.js")
        if "flask" in str(deps).lower() or "django" in str(deps).lower():
            analysis["frameworks"].append("Flask" if "flask" in str(deps).lower() else "Django")
        if "fastapi" in str(deps).lower():
            analysis["frameworks"].append("FastAPI")
        
        # Database detection
        if any(db in str(deps).lower() for db in ["mongodb", "mongoose"]):
            analysis["tech_stack"].append("MongoDB")
        if any(db in str(deps).lower() for db in ["postgresql", "psycopg", "pg"]):
            analysis["tech_stack"].append("PostgreSQL")
        if "mysql" in str(deps).lower():
            analysis["tech_stack"].append("MySQL")
        if "redis" in str(deps).lower():
            analysis["tech_stack"].append("Redis")
        
        # Structure insights
        if "api" in dirs or "backend" in dirs:
            analysis["structure_insights"].append("Backend API service")
        if "frontend" in dirs or "client" in dirs or "ui" in dirs:
            analysis["structure_insights"].append("Frontend application")
        if "docs" in dirs or "documentation" in dirs:
            analysis["structure_insights"].append("Well-documented project")
        if "tests" in dirs or "test" in dirs:
            analysis["structure_insights"].append("Includes test suite")
        if "scripts" in dirs or "bin" in dirs:
            analysis["structure_insights"].append("Automation scripts included")
        if "config" in dirs or "configs" in dirs:
            analysis["structure_insights"].append("Configurable application")
        
    except Exception as e:
        print(f"Error analyzing repository structure: {e}")
    
    return analysis

def fetch_additional_repo_signals(owner: str, repo: str, user_token: str = None) -> str:
    """Collect extra signals to help infer problem_statement and future_scope when README is sparse."""
    headers = {
        "Accept": "application/vnd.github+json",
        "User-Agent": "PitchPerfectAI-saas",
    }
    # Use user's GitHub token first (for private repos), fallback to global token
    token = user_token or GITHUB_TOKEN
    if token:
        headers["Authorization"] = f"Bearer {token}"

    parts = []

    # Repo metadata
    try:
        repo_api = f"https://api.github.com/repos/{owner}/{repo}"
        r = requests.get(repo_api, headers=headers, timeout=12)
        if r.status_code == 200:
            j = r.json()
            desc = j.get("description") or ""
            topics = j.get("topics") or []
            if desc:
                parts.append(f"Repo Description: {desc}")
            if topics:
                parts.append("Topics: " + ", ".join(topics))
    except requests.RequestException:
        pass

    # Languages
    try:
        langs_api = f"https://api.github.com/repos/{owner}/{repo}/languages"
        r = requests.get(langs_api, headers=headers, timeout=12)
        if r.status_code == 200:
            langs = r.json() or {}
            if isinstance(langs, dict) and langs:
                top = sorted(langs.items(), key=lambda kv: kv[1], reverse=True)[:6]
                lang_list = ", ".join([f"{k} ({v})" for k, v in top])
                parts.append("Languages (bytes): " + lang_list)
    except requests.RequestException:
        pass

    return "\n\n".join(parts).strip()

def build_gemini_prompt_from_code(repo_analysis: dict, repo_name: str, extra_context: str = None) -> str:
    """Build AI prompt for pitch deck generation based on code analysis instead of README."""
    
    # Extract key information from analysis
    project_type = repo_analysis.get("project_type", "Unknown")
    tech_stack = repo_analysis.get("tech_stack", [])
    main_language = repo_analysis.get("main_language", "Unknown")
    frameworks = repo_analysis.get("frameworks", [])
    features = repo_analysis.get("features", [])
    structure_insights = repo_analysis.get("structure_insights", [])
    dependencies = repo_analysis.get("dependencies", {})
    file_count = repo_analysis.get("file_count", 0)
    directories = repo_analysis.get("directory_structure", [])
    
    # Build comprehensive analysis summary
    analysis_summary = f"""
PROJECT ANALYSIS SUMMARY:
Repository Name: {repo_name}
Project Type: {project_type}
Main Language: {main_language}
File Count: {file_count}

TECHNOLOGY STACK:
{', '.join(tech_stack[:10]) if tech_stack else 'Not detected'}

FRAMEWORKS & LIBRARIES:
{', '.join(frameworks) if frameworks else 'Standard libraries'}

DIRECTORY STRUCTURE:
{', '.join(directories[:15]) if directories else 'Simple structure'}

DETECTED FEATURES:
{', '.join(features) if features else 'Core functionality'}

ARCHITECTURAL INSIGHTS:
{', '.join(structure_insights) if structure_insights else 'Standard application structure'}

DEPENDENCIES:
{str(dependencies) if dependencies else 'No dependency information available'}
"""

    parts = [
        "You are a specialized AI assistant for analyzing GitHub repositories and generating project proposals. ",
        "Your task is to analyze the provided code structure and repository information to generate a comprehensive project proposal in JSON format. ",
        "Since no README is available, you must infer the project's purpose, functionality, and value proposition from the code analysis. ",
        "Be creative but realistic in your interpretations based on the technical evidence provided.\n\n",
        
        "Instructions:\n\n",
        "1. Format: Your entire response must be a single, valid JSON object. No extra text, no markdown outside the JSON.\n",
        "2. Structure: The JSON object must have the following keys:\n",
        "   - project_title: (string) An engaging title for the project based on repo name and analysis.\n",
        "   - problem_statement: (string) Infer what problem this project likely solves based on its tech stack and structure (3-5 sentences).\n",
        "   - solution_overview: (string) Describe how the project addresses the problem based on its architecture (4-6 sentences).\n",
        "   - key_features: (string) List main features inferred from code structure and dependencies (6-10 bullet points).\n",
        "   - target_audience: (string) Who would likely use this project based on its technical nature (2-4 sentences).\n",
        "   - technology_stack: (string) Detailed explanation of the detected technologies and their purposes (4-10 elements).\n",
        "   - future_scope: (string) Potential enhancements and features that could be added based on current foundation (8-12 concrete items).\n\n",
        
        "Analysis Guidelines:\n",
        "- Use the project name and technical stack to infer the domain and purpose\n",
        "- Consider the architectural patterns evident in the directory structure\n",
        "- Infer user needs based on the frameworks and libraries used\n",
        "- Suggest realistic future features that align with the current tech stack\n",
        "- Make educated assumptions about the project's goals based on technical evidence\n\n",
        
        "Repository Analysis Data:\n\n```\n",
        analysis_summary,
        "\n```\n\n",
        ("Additional Repository Context:\n\n```\n" + extra_context.strip() + "\n```\n\n" if extra_context and extra_context.strip() else ""),
        "Generated JSON Output:",
    ]
    return "".join(parts)

def build_gemini_prompt(readme_content: str, extra_context: str = None) -> str:
    parts = [
        "You are a specialized AI assistant for a web application. Your task is to analyze the provided GitHub README content and generate a project proposal in a strict JSON format. ",
        "This JSON will be used to populate a web UI. The proposal must be based only on the information in the README. ",
        "If any information is missing for a section, set the value to 'Not available in README.'\n\n",
        "Instructions:\n\n",
        "1. Format: Your entire response must be a single, valid JSON object. No extra text, no markdown outside the JSON.\n",
        "2. Structure: The JSON object must have the following keys:\n",
        "   - title: (string) The project's title.\n",
        "   - introduction: (string) A brief overview (6-10 sentences, ~120-250 words).\n",
        "   - problem_statement: (string) The problem the project solves (6-10 sentences, ~120-250 words).\n",
        "   - solution: (string) The proposed solution (8-14 sentences, ~160-350 words).\n",
        "   - target_audience: (string) Who the project is for (2-4 sentences).\n",
        "   - technology_stack: (string) The tech used with explanations (4-10 elements).\n",
        "   - future_scope: (string) Potential future features (8-12 concrete items).\n\n",
        "README Content to Analyze:\n\n```\n",
        readme_content,
        "\n```\n\n",
        ("Additional Repository Signals:\n\n```\n" + extra_context.strip() + "\n```\n\n" if extra_context and extra_context.strip() else ""),
        "Generated JSON Output:",
    ]
    return "".join(parts)

def call_huggingface_free(readme_content: str, extra_context: str = None) -> dict:
    """Use Hugging Face free inference API for pitch deck generation"""
    import requests
    
    # Use a more capable free model for generation
    models_to_try = [
        "mistralai/Mixtral-8x7B-Instruct-v0.1",  # Very capable, free
        "meta-llama/Llama-2-70b-chat-hf",         # Good for structured output
        "tiiuae/falcon-180B-chat",                # Large, capable model
    ]
    
    # Build a simplified prompt for Hugging Face
    prompt = f"""Generate a pitch deck in JSON format for this project:

README Content:
{readme_content[:2000]}

Generate a JSON object with these fields:
- title: Project name
- introduction: Brief intro (2-3 sentences)
- problem_statement: Problem being solved (3-4 sentences)
- solution_overview: How the project solves it (3-4 sentences)
- key_features: Main features (bullet points)
- target_audience: Who it's for (2-3 sentences)
- technology_stack: Technologies used
- future_scope: Future plans (2-3 sentences)

Return ONLY valid JSON, no markdown formatting."""

    for model_url in models_to_try:
        try:
            print(f"DEBUG: Trying Hugging Face model: {model_url}")
            API_URL = f"https://api-inference.huggingface.co/models/{model_url}"
            
            response = requests.post(
                API_URL,
                headers={"Content-Type": "application/json"},
                json={"inputs": prompt[:1500], "parameters": {"max_new_tokens": 1000}},
                timeout=30
            )
            
            if response.status_code == 200:
                result = response.json()
                
                # Extract generated text
                generated_text = ""
                if isinstance(result, list) and len(result) > 0:
                    generated_text = result[0].get('generated_text', '')
                elif isinstance(result, dict):
                    generated_text = result.get('generated_text', result.get('text', ''))
                
                if generated_text and len(generated_text) > 100:
                    print(f"DEBUG: Hugging Face {model_url} success! Response length: {len(generated_text)}")
                    
                    # Try to parse JSON from the response
                    try:
                        # Clean up the response
                        json_start = generated_text.find('{')
                        json_end = generated_text.rfind('}') + 1
                        if json_start >= 0 and json_end > json_start:
                            json_str = generated_text[json_start:json_end]
                            parsed = json.loads(json_str)
                            return parsed
                    except:
                        pass
            
            print(f"DEBUG: Hugging Face {model_url} failed with status: {response.status_code}")
            
        except Exception as e:
            print(f"DEBUG: Hugging Face {model_url} error: {str(e)}")
            continue
    
    # If Hugging Face fails, return a basic template
    print("DEBUG: All Hugging Face models failed, using template")
    return generate_basic_template(readme_content)

def generate_basic_template(readme_content: str) -> dict:
    """Generate a basic pitch deck template from README content"""
    # Extract project name from first line or use default
    lines = readme_content.split('\n')
    title = lines[0].strip('#').strip() if lines else "Project Pitch Deck"
    
    return {
        "title": title,
        "introduction": f"This is an innovative project that aims to solve real-world problems. Based on the codebase analysis, this project shows strong technical implementation and clear value proposition.",
        "problem_statement": "Many users face challenges that require efficient, scalable solutions. Current alternatives are often complex, expensive, or lack key features that users need.",
        "solution_overview": f"{title} addresses these challenges through a well-architected solution that combines modern technology with user-centric design. The implementation focuses on reliability, performance, and ease of use.",
        "key_features": "• Modern, scalable architecture\n• User-friendly interface\n• Robust error handling\n• Comprehensive documentation\n• Active development and maintenance",
        "target_audience": "This solution is designed for developers, businesses, and organizations looking for reliable, efficient tools. It serves both technical and non-technical users who need powerful yet accessible solutions.",
        "technology_stack": "Built with modern, industry-standard technologies ensuring reliability, maintainability, and scalability. The stack is chosen for optimal performance and developer experience.",
        "future_scope": "Future development will focus on expanding features, improving performance, and incorporating user feedback. Plans include enhanced integrations, additional customization options, and continued optimization."
    }

def call_gemini(readme_content: str, extra_context: str = None) -> dict:
    """
    REPLACED: Now uses Hugging Face instead of Gemini
    Kept function name for backward compatibility
    """
    print("DEBUG: Using Hugging Face for pitch generation (Gemini replaced)")
    return generate_pitch_json(readme_content, extra_context)

def parse_gemini_json(text: str) -> dict:
    """Parse JSON response from Gemini"""
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        m = re.search(r"\{[\s\S]*\}", text)
        if not m:
            raise
        parsed = json.loads(m.group(0))

    return parsed

@api_bp.route('/fetch-readme', methods=['POST'])
@login_required
def api_fetch_readme():
    data = request.get_json(silent=True) or {}
    repo_url = (data.get("repo_url") or "").strip()

    owner, repo = parse_github_repo(repo_url)
    if not owner:
        return jsonify({"validation_status": False, "error": "Invalid GitHub repository URL."}), 400

    content, source = try_fetch_readme_raw(owner, repo)
    
    # If that fails, try GitHub API with user's token
    if not content:
        user_token = current_user.github_token if current_user.is_authenticated else None
        content, source = try_fetch_readme_api(owner, repo, user_token)

    if not content:
        return jsonify({"validation_status": False, "error": "README not found in repository."}), 404

    return jsonify({
        "validation_status": True,
        "readme_content": content,
        "source": source,
    })

@api_bp.route('/generate', methods=['POST'])
@login_required
def api_generate():
    # Check if user has tokens or is waitlisted with survey completed
    from dashboard import _has_generation_access
    if not _has_generation_access(current_user):
        return jsonify({'error': 'No tokens remaining or access denied'}), 403

    # Handle both form data and JSON data
    if request.content_type and 'multipart/form-data' in request.content_type:
        # Form data with potential file upload
        project_description = request.form.get('project_description', '').strip()
        project_file = request.files.get('project_file')
        generation_type = request.form.get('generation_type', 'investor')
        
        # Get user preferences for slide count and other settings
        try:
            preferences = session.get('user_preferences', {})
            slide_count = int(preferences.get('slide_count', '12'))
            include_financials = preferences.get('include_financials', True)
            include_competition = preferences.get('include_competition', True)
        except:
            # Fallback to defaults if session access fails
            slide_count = 12
            include_financials = True
            include_competition = True
        
        content = ""
        if project_file:
            # Process uploaded file
            try:
                if project_file.filename.endswith(('.txt', '.md')):
                    content = project_file.read().decode('utf-8')
                elif project_file.filename.endswith('.pdf'):
                    # For PDF files, you'd need a PDF parser like PyPDF2
                    return jsonify({"error": "PDF processing not yet implemented"}), 400
                elif project_file.filename.endswith(('.doc', '.docx')):
                    # For Word docs, you'd need python-docx
                    return jsonify({"error": "Word document processing not yet implemented"}), 400
                else:
                    return jsonify({"error": "Unsupported file format"}), 400
            except Exception as e:
                return jsonify({"error": f"Failed to process file: {str(e)}"}), 400
        elif project_description:
            content = project_description
        else:
            return jsonify({"error": "Please provide either a project description or upload a document"}), 400
            
        # Set project details from description
        title = "Custom Project"
        repo_owner = current_user.username
        repo_name = "custom-project"
        repo_url = ""
        
    else:
        # Legacy JSON data for GitHub repos (keep for backward compatibility)
        data = request.get_json(silent=True) or {}
        repo_url = (data.get("repo_url") or "").strip()

        owner, repo = parse_github_repo(repo_url)
        if not owner:
            return jsonify({"error": "Invalid GitHub repository URL."}), 400

        # Try raw.githubusercontent.com first
        content, source = try_fetch_readme_raw(owner, repo)
        
        # If that fails, try GitHub API with user's token
        if not content:
            user_token = current_user.github_token if current_user.is_authenticated else None
            content, _ = try_fetch_readme_api(owner, repo, user_token)

        if not content:
            # No README found - use code analysis instead
            print(f"DEBUG: No README found, analyzing repository structure for {owner}/{repo}")
            user_token = current_user.github_token if current_user.is_authenticated else None
            repo_analysis = analyze_repository_structure(owner, repo, user_token)
            
            if not repo_analysis.get("tech_stack") and repo_analysis.get("main_language") == "Unknown":
                return jsonify({"error": "Unable to analyze repository structure. Repository may be empty or inaccessible."}), 404
            
            # Generate pitch deck from code analysis using Hugging Face
            try:
                print(f"DEBUG: Using Hugging Face for code analysis pitch generation")
                # Create a summary from repo analysis
                code_summary = f"""
Project: {repo}
Main Language: {repo_analysis.get('main_language', 'Unknown')}
Tech Stack: {', '.join(repo_analysis.get('tech_stack', []))}
File Count: {repo_analysis.get('file_count', 0)}
Directory Structure: {', '.join(repo_analysis.get('directory_structure', [])[:10])}
"""
                
                # Use Hugging Face to generate pitch from code analysis
                result = generate_pitch_json(code_summary, None)
                
                return jsonify({
                    "validation_status": True,
                    "analysis_result": result,
                    "source": "code_analysis",
                    "repo_analysis": repo_analysis
                })
                
            except Exception as e:
                return jsonify({"error": f"Error generating pitch deck from code analysis: {str(e)}"}), 500
            
        title = repo
        repo_owner = owner
        repo_name = repo

    try:
        # REPLACED: Now using Hugging Face instead of Gemini
        print("DEBUG: Using Hugging Face for pitch generation")
        
        # For custom projects, we don't fetch additional repo signals
        if 'multipart/form-data' in request.content_type:
            extra = {}
        else:
            user_token = current_user.github_token if current_user.is_authenticated else None
            extra = fetch_additional_repo_signals(repo_owner, repo_name, user_token)
        
        # Build dynamic slide structure based on preferences
        slides = [
            "**Slide 1: Title Slide**\n   - Project name and tagline\n   - Team/creator information\n   - Date",
            "**Slide 2: Problem**\n   - What problem does this project solve?\n   - Pain points and market gaps\n   - Why this matters now",
            "**Slide 3: Solution**\n   - How does this project address the problem?\n   - Key features and functionality\n   - Unique value proposition",
            "**Slide 4: Market Opportunity**\n   - Target market size\n   - User personas and segments\n   - Market trends and timing",
            "**Slide 5: Product Demo**\n   - Key features walkthrough\n   - User experience highlights\n   - Technical capabilities",
            "**Slide 6: Technology Stack**\n   - Architecture overview\n   - Key technologies used\n   - Technical advantages",
            "**Slide 7: Traction & Metrics**\n   - Current usage/adoption\n   - Key performance indicators\n   - Growth metrics"
        ]
        
        # Add competition slide if enabled
        if include_competition:
            slides.append("**Slide 8: Competition**\n   - Competitive landscape\n   - Competitive advantages\n   - Differentiation strategy")
        
        # Add business model and financials
        slides.append("**Slide {}: Business Model**\n   - Revenue streams\n   - Monetization strategy\n   - Pricing approach".format(len(slides) + 1))
        
        if include_financials:
            slides.append("**Slide {}: Financials**\n   - Revenue projections\n   - Cost structure\n   - Funding requirements".format(len(slides) + 1))
        
        # Add remaining core slides
        slides.extend([
            "**Slide {}: Team**\n   - Key team members\n   - Relevant experience\n   - Advisory board".format(len(slides) + 1),
            "**Slide {}: Roadmap**\n   - Future features and milestones\n   - Development timeline\n   - Strategic vision".format(len(slides) + 1),
            "**Slide {}: Ask & Next Steps**\n   - What you're seeking (funding, partnerships, users)\n   - Use of funds/resources\n   - Call to action".format(len(slides) + 1)
        ])
        
        # Add additional slides if user wants more than the core set
        core_slides = len(slides)
        if slide_count > core_slides:
            additional_slides = slide_count - core_slides
            for i in range(additional_slides):
                slide_num = core_slides + i + 1
                if i == 0:
                    slides.append(f"**Slide {slide_num}: Market Analysis**\n   - Detailed market research\n   - Customer segments\n   - Market size validation")
                elif i == 1:
                    slides.append(f"**Slide {slide_num}: Product Roadmap**\n   - Feature development timeline\n   - Version releases\n   - Long-term vision")
                elif i == 2:
                    slides.append(f"**Slide {slide_num}: Risk Analysis**\n   - Potential challenges\n   - Mitigation strategies\n   - Contingency plans")
                elif i == 3:
                    slides.append(f"**Slide {slide_num}: Partnership Strategy**\n   - Strategic partnerships\n   - Distribution channels\n   - Ecosystem integration")
                else:
                    slides.append(f"**Slide {slide_num}: Additional Details**\n   - Supporting information\n   - Technical specifications\n   - Implementation details")
        
        # Limit to requested slide count
        slides = slides[:slide_count]
        
        slide_structure = "\n\n".join(slides)
        
        # Extract project details for market research
        project_name = title
        technology_stack = "Unknown"
        target_audience = "General"
        
        # Try to extract tech stack and audience from content
        content_lower = content.lower()
        if any(tech in content_lower for tech in ["javascript", "js", "node", "react", "vue", "angular"]):
            technology_stack = "JavaScript/Web Development"
        elif any(tech in content_lower for tech in ["python", "django", "flask", "fastapi"]):
            technology_stack = "Python"
        elif any(tech in content_lower for tech in ["mobile", "ios", "android", "flutter", "react native"]):
            technology_stack = "Mobile Development"
        elif any(tech in content_lower for tech in ["ai", "ml", "machine learning", "neural", "gpt"]):
            technology_stack = "AI/Machine Learning"
        
        if any(audience in content_lower for audience in ["business", "enterprise", "company"]):
            target_audience = "Businesses"
        elif any(audience in content_lower for audience in ["developer", "programmer", "coder"]):
            target_audience = "Developers"
        elif any(audience in content_lower for audience in ["startup", "entrepreneur"]):
            target_audience = "Startups"
        
        # Generate market intelligence for this project
        from dashboard import generate_market_intelligence
        market_data = generate_market_intelligence(project_name, technology_stack, target_audience, content)
        
        # Build enhanced prompt with market intelligence
        market_context = f"""
Market Intelligence:
- Market Size: {market_data.get('market_size', 'Market analysis pending')}
- Direct Competitors: {market_data.get('competitors_direct', 'Competitive analysis pending')}
- Indirect Competitors: {market_data.get('competitors_indirect', 'Alternative solutions analysis pending')}
- Industry Trends: {market_data.get('industry_trends', 'Trend analysis pending')}
- Target Market Insights: {market_data.get('target_market_insights', 'Audience analysis pending')}
- Growth Projections: {market_data.get('growth_projections', 'Growth analysis pending')}
- Customer Acquisition Cost: {market_data.get('cac_estimate', 'CAC analysis pending')}
- Customer Lifetime Value: {market_data.get('ltv_estimate', 'LTV analysis pending')}
"""

        prompt = f"""Generate a comprehensive pitch deck for the following project using the provided market intelligence:

Project: {title}
{f'Repository: {repo_owner}/{repo_name}' if repo_url else 'Custom Project'}

Project Content:\n{content}\n\n

{market_context}

Please generate a detailed pitch deck with exactly {slide_count} slides as specified below:

{slide_structure}

CRITICAL ANTI-PLACEHOLDER RULES:
1. NO placeholder text anywhere - no [Company Name], [X%], [Insert Details], [TBD], etc.
2. For PROBLEM STATEMENT: Use specific, real-world problems with concrete examples and statistics
3. For SOLUTION: Describe actual features and capabilities based on the project content
4. For MARKET OPPORTUNITY: Use the provided market intelligence data with specific figures
5. For COMPETITION: Name real competitors with actual pricing and feature comparisons
6. For BUSINESS MODEL: Propose realistic revenue streams based on similar successful projects
7. For FINANCIALS: Use industry-standard metrics and realistic projections
8. For TEAM: If team info not provided, focus on required expertise and roles needed
9. For ROADMAP: Create realistic development milestones based on project scope
10. For NEXT STEPS: Specify concrete actions, funding amounts, and timelines
11. For VISION: Articulate specific long-term goals and market impact

CONTENT SOURCING REQUIREMENTS:
- Extract all details from the project content provided
- Use market intelligence data for competitive and financial information
- Apply industry knowledge to fill gaps with realistic, specific information
- Never use generic terms or ask readers to "insert" information
- Every slide must be complete and presentation-ready

Make each slide investor-ready with specific, actionable content that tells a compelling story.
"""
        
        # Generate pitch deck using Hugging Face
        print("DEBUG: Calling Hugging Face for pitch generation...")
        result = generate_pitch_json(content, market_context)
        
        # Format the result as pitch deck content
        pitch_deck_content = f"""# {result.get('title', title)}

## Introduction
{result.get('introduction', '')}

## Problem Statement
{result.get('problem_statement', '')}

## Solution Overview
{result.get('solution_overview', '')}

## Key Features
{result.get('key_features', '')}

## Target Audience
{result.get('target_audience', '')}

## Technology Stack
{result.get('technology_stack', '')}

## Future Scope
{result.get('future_scope', '')}
"""
        
        # Save pitch deck to project
        project = Project(
            title=title,
            repo_url=repo_url,
            repo_owner=repo_owner,
            repo_name=repo_name,
            pitch_deck={
                'content': pitch_deck_content,
                'generated_at': datetime.utcnow().isoformat(),
                'version': '1.0'
            },
            user_id=current_user.id
        )
        
        project.save()
        
        # No token usage for generation - tokens are only for exports
        
        return jsonify({
            'success': True,
            'pitch_deck': project.pitch_deck,
            'project_id': project.id,
            'tokens_remaining': current_user.tokens
        })
    
    except Exception as e:
        return jsonify({'error': f'Failed to generate pitch deck: {str(e)}'}), 500

def _pitchy_project(project_id):
    """Load a project owned by the current user, or None."""
    from firebase_models import Project
    project = Project.get(project_id) if project_id else None
    if not project or project.user_id != current_user.id:
        return None
    return project


def _pitch_content(project):
    return (project.pitch_deck or {}).get('content', '') if project else ''


def _pitch_slide_count(project):
    """Number of content slides for a fixed-length pitch (its section count), else None."""
    if not project or not (project.pitch_deck or {}).get('slide_count'):
        return None
    import pitchy
    return len(pitchy.split_sections(_pitch_content(project))) or None


def _pitchy_error(e, action):
    """Log the raw failure; return a friendly message (raw details for admins)."""
    from huggingface_client import friendly_error
    print(f"ERROR: Pitchy {action} failed for user {current_user.id}: {e}")
    payload = {'error': friendly_error(e, 'talking to Pitchy')}
    if getattr(current_user, 'is_admin', False):
        payload['details'] = str(e)
    return jsonify(payload), 503


@api_bp.route('/pitchy-review', methods=['POST'])
@login_required
def pitchy_review():
    """Score each pitch section and suggest what to strengthen."""
    import pitchy
    data = request.get_json(silent=True) or {}
    project = _pitchy_project(data.get('project_id'))
    if not project:
        return jsonify({'error': 'Project not found'}), 404
    content = _pitch_content(project)
    if not content.strip():
        return jsonify({'error': 'This project has no pitch deck to review yet.'}), 400
    try:
        return jsonify({'success': True, **pitchy.review_pitch(project.title, content, _pitch_slide_count(project))})
    except Exception as e:
        return _pitchy_error(e, 'review')


@api_bp.route('/pitchy-chat', methods=['POST'])
@login_required
def pitchy_chat():
    """Chat with Pitchy about the project's pitch; may propose a section rewrite."""
    import pitchy
    data = request.get_json(silent=True) or {}
    message = (data.get('message') or '').strip()
    if not message:
        return jsonify({'error': 'Message is required'}), 400
    project = _pitchy_project(data.get('project_id'))
    if not project:
        return jsonify({'error': 'Project not found'}), 404
    history = data.get('history') if isinstance(data.get('history'), list) else []
    try:
        result = pitchy.chat(project.title, _pitch_content(project), history, message, _pitch_slide_count(project))
        return jsonify({'success': True, **result})
    except Exception as e:
        return _pitchy_error(e, 'chat')


@api_bp.route('/pitchy-apply', methods=['POST'])
@login_required
def pitchy_apply():
    """Apply a Pitchy section rewrite to the pitch (previous version kept for undo)."""
    import pitchy
    from dashboard import save_pitch_version
    data = request.get_json(silent=True) or {}
    project = _pitchy_project(data.get('project_id'))
    section = (data.get('section') or '').strip()
    new_body = (data.get('content') or '').strip()
    if not project:
        return jsonify({'error': 'Project not found'}), 404
    if not section or not new_body:
        return jsonify({'error': 'Section and content are required'}), 400
    if _pitch_slide_count(project) and not pitchy.section_exists(_pitch_content(project), section):
        return jsonify({'error': f'This deck is fixed at {_pitch_slide_count(project)} slides, so new sections '
                                 'can\'t be added. Ask Pitchy to fold this into an existing section.'}), 400
    try:
        updated, heading = pitchy.replace_section(_pitch_content(project), section[:80], new_body[:8000])
        save_pitch_version(project, updated, 'pitchy_ai')
        return jsonify({
            'success': True,
            'section': heading,
            'html': current_app.jinja_env.filters['format_markdown'](updated)
        })
    except Exception as e:
        print(f"ERROR: Pitchy apply failed for user {current_user.id}: {e}")
        return jsonify({'error': 'Could not apply the change. Please try again.'}), 500


def _save_generated_pitch(title, pitch_content, slide_count, repo=None):
    """Save a freshly generated pitch as a project and build the JSON response."""
    from firebase_models import Project
    project = Project(
        title=title,
        repo_url=repo['url'] if repo else "",
        repo_owner=repo['owner'] if repo else current_user.username,
        repo_name=repo['name'] if repo else title.lower().replace(' ', '-'),
        pitch_deck={
            'content': pitch_content,
            'generated_at': datetime.utcnow().isoformat(),
            'version': '2.0',
            'generation_type': 'github_repo' if repo else 'deep_pipeline',
            'slide_count': slide_count
        },
        user_id=current_user.id
    )
    project.save()
    return jsonify({
        'success': True,
        'pitch': pitch_content,
        # Same rendering as the project page, so the preview matches it
        'html': current_app.jinja_env.filters['format_markdown'](pitch_content),
        'title': title,
        'project_id': project.id,
        'tokens_remaining': current_user.tokens
    })


def _generation_error(e):
    from huggingface_client import friendly_error
    print(f"ERROR: Pitch generation failed for user {current_user.id}: {str(e)}")
    payload = {'error': friendly_error(e)}
    if getattr(current_user, 'is_admin', False):
        payload['details'] = str(e)
    return jsonify(payload), 500


@api_bp.route('/generate-deep-pitch', methods=['POST'])
@login_required
def generate_deep_pitch():
    """
    Generate an investor pitch from a name and description.
    Expects JSON: { "name": "Startup Name", "description": "Brief description" }
    """
    from dashboard import _has_generation_access
    if not _has_generation_access(current_user):
        return jsonify({'error': 'No tokens remaining or access denied'}), 403

    data = request.get_json(silent=True) or {}
    startup_name = (data.get('name') or '').strip()
    startup_description = (data.get('description') or '').strip()
    if not startup_name or not startup_description:
        return jsonify({'error': 'Both startup name and description are required'}), 400

    try:
        from slide_plan import user_slide_count
        slide_count = user_slide_count(current_user)
        pitch_content = deep_pitch_generation(startup_name, startup_description, slide_count)
        return _save_generated_pitch(startup_name, pitch_content, slide_count)
    except Exception as e:
        return _generation_error(e)


@api_bp.route('/generate-from-repo', methods=['POST'])
@login_required
def generate_from_repo():
    """
    Generate an investor pitch from a GitHub repository.
    Expects JSON: { "repo": "owner/name" or a github.com URL }
    """
    from dashboard import _has_generation_access
    from github_repo import fetch_repo_context, RepoError
    from huggingface_client import intelligent_pitch_generation
    from slide_plan import user_slide_count

    if not _has_generation_access(current_user):
        return jsonify({'error': 'No tokens remaining or access denied'}), 403

    data = request.get_json(silent=True) or {}
    try:
        repo = fetch_repo_context(data.get('repo'), getattr(current_user, 'github_token', None))
    except RepoError as e:
        return jsonify({'error': str(e)}), 400

    try:
        slide_count = user_slide_count(current_user)
        # Leave room for the GitHub calls already made within Vercel's 60s limit
        pitch_content = intelligent_pitch_generation(repo['title'], repo['context'], slide_count,
                                                     from_repo=True, budget_seconds=42)
        return _save_generated_pitch(repo['title'], pitch_content, slide_count, repo)
    except Exception as e:
        return _generation_error(e)


@api_bp.route('/projects/<project_id>', methods=['GET'])
@login_required
def get_project(project_id):
    project = Project.get(project_id)
    if not project or project.user_id != current_user.id:
        return jsonify({'error': 'Project not found'}), 404
    
    return jsonify({
        'id': project.id,
        'title': project.title,
        'repo_url': project.repo_url,
        'repo_owner': project.repo_owner,
        'repo_name': project.repo_name,
        'pitch_deck': project.pitch_deck,
        'user_id': project.user_id,
        'created_at': project.created_at.isoformat() if project.created_at else None,
        'updated_at': project.updated_at.isoformat() if project.updated_at else None
    })

@api_bp.route('/projects/<project_id>/export/docx', methods=['GET'])
@login_required
def export_project_docx(project_id):
    # Check if user has tokens for export
    if not current_user.can_export_files():
        return redirect(url_for('dashboard.out_of_tokens'))
    
    project = Project.get(project_id)
    if not project or project.user_id != current_user.id:
        return jsonify({'error': 'Project not found'}), 404
    
    # Create DOCX document
    doc = Document()
    
    # Title
    title = doc.add_heading(project.title, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Repository info
    doc.add_paragraph()
    doc.add_paragraph(f"Repository: {project.repo_owner}/{project.repo_name}")
    doc.add_paragraph(f"URL: {project.repo_url}")
    doc.add_paragraph(f"Generated: {project.created_at.strftime('%B %d, %Y')}")
    doc.add_paragraph()
    
    # Pitch deck content
    if project.pitch_deck and project.pitch_deck.get('content'):
        content = project.pitch_deck.get('content', '')
        
        # Split content by slide markers or use as single content
        if '**Slide' in content:
            # Parse structured slide content
            slides = content.split('**Slide')[1:]  # Skip empty first element
            for slide in slides:
                lines = slide.strip().split('\n')
                if lines:
                    # First line is the slide title
                    slide_title = lines[0].replace(':', '').strip()
                    doc.add_heading(f"Slide {slide_title}", level=1)
                    
                    # Rest is content
                    slide_content = '\n'.join(lines[1:]).strip()
                    if slide_content:
                        doc.add_paragraph(slide_content)
                    doc.add_paragraph()
        else:
            # Use content as-is
            doc.add_paragraph(content)
    
    # Save to buffer
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    # Deduct token for export
    current_user.use_token()
    
    filename = f"{project.title.replace(' ', '_')}_pitch_deck.docx"
    return send_file(buffer, as_attachment=True, download_name=filename, 
                    mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@api_bp.route('/projects/<project_id>/export/pptx', methods=['GET'])
@login_required
def export_project_pptx(project_id):
    # Check if user has tokens for export
    if not current_user.can_export_files():
        return redirect(url_for('dashboard.out_of_tokens'))
    
    project = Project.get(project_id)
    if not project or project.user_id != current_user.id:
        return jsonify({'error': 'Project not found'}), 404
    
    # Build the branded deck from the pitch Markdown
    from deck_builder import build_pitch_deck
    content = (project.pitch_deck or {}).get('content', '')
    theme = {'1': 'coral', '2': 'green', '3': 'red'}.get(request.args.get('template', '1'), 'coral')
    # Use the length the pitch was written for; older pitches follow the current preference
    from slide_plan import user_slide_count
    stored = (project.pitch_deck or {}).get('slide_count')
    slide_count = int(stored) if stored else user_slide_count(current_user)
    buffer = build_pitch_deck(project.title, content, theme, slide_count)
    
    # Deduct token for export
    current_user.use_token()
    
    filename = f"{project.title.replace(' ', '_')}_pitch_deck.pptx"
    return send_file(buffer, as_attachment=True, download_name=filename, 
                    mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation')

@api_bp.route('/generate-documentation/<project_id>')
@login_required
def generate_documentation(project_id):
    """Generate technical documentation for a project"""
    # Check if user has generation access
    from dashboard import _has_generation_access
    if not _has_generation_access(current_user):
        return jsonify({'error': 'No tokens remaining or access denied'}), 403
    
    project = Project.get(project_id)
    
    if not project or project.user_id != current_user.id:
        return jsonify({'error': 'Project not found or access denied'}), 404
    
    try:
        # REPLACED: Now using Hugging Face for documentation generation
        print("DEBUG: Using Hugging Face for documentation generation")
        
        # Get repository content for context
        repo_content = ""
        if project.repo_owner and project.repo_name:
            readme_content, _ = try_fetch_readme_raw(project.repo_owner, project.repo_name)
            if not readme_content:
                user_token = current_user.github_token if current_user.is_authenticated else None
                readme_content, _ = try_fetch_readme_api(project.repo_owner, project.repo_name, user_token)
            if readme_content:
                repo_content = f"README Content:\n{readme_content}\n\n"
        
        # Create documentation generation prompt
        prompt = f"""Generate technical documentation for this project:

Project: {project.title}
Repository: {project.repo_owner}/{project.repo_name if project.repo_name else 'N/A'}

{repo_content[:1000]}

Include: Architecture, API docs, Installation, Development guide, Deployment, Troubleshooting.
Format as markdown.
"""
        
        # Generate documentation using Hugging Face
        documentation_content = refine_pitch_content("", prompt)
        
        if not documentation_content:
            documentation_content = f"""# {project.title} Documentation

## Overview
Technical documentation for {project.title}.

## Installation
```bash
# Clone the repository
git clone {project.repo_url if project.repo_url else 'repository-url'}

# Install dependencies
npm install  # or pip install -r requirements.txt
```

## Usage
Please refer to the README for detailed usage instructions.

## Development
Contributions are welcome! Please follow the project's coding standards.

## License
See LICENSE file for details.
"""
        
        # Save documentation to project
        project.documentation = {
            'content': documentation_content,
            'generated_at': datetime.utcnow().isoformat(),
            'version': '1.0'
        }
        
        # No token usage for generation - tokens are only for exports
        project.save()
        
        return jsonify({
            'success': True,
            'documentation': project.documentation,
            'tokens_remaining': current_user.tokens
        })
        
    except Exception as e:
        return jsonify({'error': f'Documentation generation failed: {str(e)}'}), 500

@api_bp.route('/generate-user-guide/<project_id>')
@login_required
def generate_user_guide(project_id):
    """Generate user guide for a project"""
    # Check if user has generation access
    from dashboard import _has_generation_access
    if not _has_generation_access(current_user):
        return jsonify({'error': 'No tokens remaining or access denied'}), 403
    
    project = Project.get(project_id)
    
    if not project or project.user_id != current_user.id:
        return jsonify({'error': 'Project not found or access denied'}), 404
    
    try:
        # REPLACED: Now using Hugging Face for user guide generation
        print("DEBUG: Using Hugging Face for user guide generation")
        
        # Get repository content for context
        repo_content = ""
        if project.repo_owner and project.repo_name:
            readme_content, _ = try_fetch_readme_raw(project.repo_owner, project.repo_name)
            if not readme_content:
                user_token = current_user.github_token if current_user.is_authenticated else None
                readme_content, _ = try_fetch_readme_api(project.repo_owner, project.repo_name, user_token)
            if readme_content:
                repo_content = f"README Content:\n{readme_content}\n\n"
        
        # Create user guide generation prompt
        prompt = f"""Generate a user guide for this project:

Project: {project.title}
Repository: {project.repo_owner}/{project.repo_name if project.repo_name else 'N/A'}

{repo_content[:1000]}

Include: Getting Started, Installation, User Interface, Tutorials, FAQ, Troubleshooting.
Format as markdown.
   - Tips and best practices

5. **Features Reference**
   - Complete feature list with descriptions
   - Settings and customization options
   - Advanced features

6. **Troubleshooting & FAQ**
   - Common user issues and solutions
   - Frequently asked questions
   - Where to get help

7. **Tips & Best Practices**
   - How to get the most out of the application
   - Performance tips
   - Security considerations for users

Write in a friendly, accessible tone suitable for end users. Use clear headings, bullet points, and step-by-step instructions.
"""
        
        # Generate user guide using Hugging Face
        user_guide_content = refine_pitch_content("", prompt)
        
        if not user_guide_content:
            user_guide_content = f"""# {project.title} User Guide

## Getting Started
Welcome to {project.title}! This guide will help you get up and running quickly.

## Installation
```bash
# Clone the repository
git clone {project.repo_url if project.repo_url else 'repository-url'}

# Install dependencies
npm install  # or pip install -r requirements.txt

# Run the application
npm start  # or python app.py
```

## Features
- Easy to use interface
- Comprehensive functionality
- Regular updates and improvements

## Support
For questions or issues, please refer to the project repository or contact the maintainers.
"""
        
        # Save user guide to project
        project.user_guide = {
            'content': user_guide_content,
            'generated_at': datetime.utcnow().isoformat(),
            'version': '1.0'
        }
        
        # No token usage for generation - tokens are only for exports
        project.save()
        
        return jsonify({
            'success': True,
            'user_guide': project.user_guide,
            'tokens_remaining': current_user.tokens
        })
        
    except Exception as e:
        return jsonify({'error': f'User guide generation failed: {str(e)}'}), 500


@api_bp.route('/projects/<project_id>/export/pdf', methods=['GET'])
@login_required
def export_project_pdf(project_id):
    # Check if user has tokens for export
    if not current_user.can_export_files():
        return redirect(url_for('dashboard.out_of_tokens'))
    
    # PDF export temporarily disabled due to reportlab build issues on Vercel
    return jsonify({
        'error': 'PDF export temporarily unavailable',
        'message': 'Please use DOCX export or copy the content manually'
    }), 503

@api_bp.route('/upgrade-account', methods=['POST'])
@login_required
def upgrade_account():
    """Upgrade user account tier"""
    data = request.get_json()
    new_tier = data.get('tier')
    
    if new_tier not in ['basic', 'pro']:
        return jsonify({'success': False, 'error': 'Invalid tier'}), 400
    
    # In a real app, you'd integrate with a payment processor here
    # For demo purposes, we'll just upgrade the account
    if current_user.upgrade_account(new_tier):
        return jsonify({
            'success': True, 
            'tier': current_user.account_tier,
            'tokens': current_user.tokens
        })
    else:
        return jsonify({'success': False, 'error': 'Upgrade failed'}), 500

@api_bp.route('/admin/change-user-plan', methods=['POST'])
@login_required
def admin_change_user_plan():
    """Admin endpoint to change user account plans"""
    if not current_user.is_admin:
        return jsonify({'success': False, 'error': 'Admin access required'}), 403
    
    data = request.get_json()
    user_id = data.get('user_id')
    new_tier = data.get('tier')
    
    if not user_id or new_tier not in ['free', 'basic', 'pro']:
        return jsonify({'success': False, 'error': 'Invalid parameters'}), 400
    
    # Get the user to modify
    from firebase_models import User
    user = User.get(user_id)
    if not user:
        return jsonify({'success': False, 'error': 'User not found'}), 404
    
    # Change the user's plan
    if user.upgrade_account(new_tier):
        return jsonify({
            'success': True,
            'user_id': user_id,
            'tier': user.account_tier,
            'tokens': user.tokens
        })
    else:
        return jsonify({'success': False, 'error': 'Plan change failed'}), 500


@api_bp.route('/admin/toggle-admin-status', methods=['POST'])
@login_required
def toggle_admin_status():
    """Toggle admin status for a user (admin only)"""
    if not current_user.is_admin:
        return jsonify({'error': 'Admin access required'}), 403
    
    data = request.get_json()
    user_id = data.get('user_id')
    is_admin = data.get('is_admin')
    
    if not user_id or is_admin is None:
        return jsonify({'success': False, 'error': 'Invalid parameters'}), 400
    
    # Get the user to modify
    from firebase_models import User
    user = User.get(user_id)
    if not user:
        return jsonify({'success': False, 'error': 'User not found'}), 404
    
    # Don't allow removing admin from self
    if user.id == current_user.id and not is_admin:
        return jsonify({'success': False, 'error': 'Cannot remove admin privileges from yourself'}), 400
    
    # Update admin status
    user.is_admin = is_admin
    if user.save():
        return jsonify({
            'success': True,
            'user_id': user_id,
            'is_admin': user.is_admin
        })
    else:
        return jsonify({'success': False, 'error': 'Admin status change failed'}), 500


@api_bp.route('/admin/whitelist-user', methods=['POST'])
@login_required
def whitelist_user():
    """Toggle whitelist status for a deleted user (admin only)"""
    if not current_user.is_admin:
        return jsonify({'error': 'Admin access required'}), 403
    
    data = request.get_json()
    user_id = data.get('user_id')
    is_whitelisted = data.get('is_whitelisted')
    
    if not user_id or is_whitelisted is None:
        return jsonify({'success': False, 'error': 'Invalid parameters'}), 400
    
    # Get the user to modify
    user = User.get(user_id)
    
    if not user:
        return jsonify({'success': False, 'error': 'User not found'}), 404
    
    # Only allow whitelisting deleted users
    if not user.is_deleted:
        return jsonify({'success': False, 'error': 'Can only whitelist deleted users'}), 400
    
    # Update whitelist status
    user.is_whitelisted = is_whitelisted
    if user.save():
        return jsonify({
            'success': True,
            'user_id': user_id,
            'is_whitelisted': user.is_whitelisted
        })
    else:
        return jsonify({'success': False, 'error': 'Whitelist status change failed'}), 500


@api_bp.route('/admin/delete-user', methods=['DELETE'])
@login_required
def delete_user():
    """Delete a user account and all associated data (admin only)"""
    if not current_user.is_admin:
        return jsonify({'error': 'Admin access required'}), 403
    
    data = request.get_json()
    user_id = data.get('user_id')
    
    if not user_id:
        return jsonify({'success': False, 'error': 'User ID required'}), 400
    
    # Don't allow deleting self
    if user_id == current_user.id:
        return jsonify({'success': False, 'error': 'Cannot delete your own account'}), 400
    
    # Get the user to delete
    from firebase_models import User, Project
    user = User.get(user_id)
    if not user:
        return jsonify({'success': False, 'error': 'User not found'}), 404
    
    try:
        # Delete all user's projects first
        projects = Project.get_by_user(user_id)
        for project in projects:
            project.delete()
        
        # Delete the user account
        if user.delete():
            return jsonify({
                'success': True,
                'message': f'User account and all associated data deleted successfully'
            })
        else:
            return jsonify({'success': False, 'error': 'User deletion failed'}), 500
            
    except Exception as e:
        return jsonify({'success': False, 'error': f'Deletion failed: {str(e)}'}), 500


@api_bp.route('/admin/permanently-delete-user', methods=['DELETE'])
@login_required
def permanently_delete_user():
    """Permanently delete a user account and all associated data from database (admin only)"""
    if not current_user.is_admin:
        return jsonify({'error': 'Admin access required'}), 403
    
    data = request.get_json()
    user_id = data.get('user_id')
    
    if not user_id:
        return jsonify({'success': False, 'error': 'User ID required'}), 400
    
    # Don't allow deleting self
    if user_id == current_user.id:
        return jsonify({'success': False, 'error': 'Cannot delete your own account'}), 400
    
    # Get the user to delete
    from firebase_models import User, Project, SurveyResponse
    user = User.get(user_id)
    if not user:
        return jsonify({'success': False, 'error': 'User not found'}), 404
    
    # Only allow permanent deletion of already soft-deleted users
    if not getattr(user, 'is_deleted', False):
        return jsonify({'success': False, 'error': 'User must be soft-deleted first'}), 400
    
    try:
        # Delete all user's projects first
        projects = Project.get_by_user(user_id)
        for project in projects:
            project.delete()
        
        # Delete all user's survey responses
        survey_responses = SurveyResponse.get_by_user(user_id)
        for response in survey_responses:
            response.delete()
        
        # Permanently delete the user account from Firestore
        from firebase_config import get_db
        db = get_db()
        if db:
            db.collection('users').document(user_id).delete()
            return jsonify({
                'success': True,
                'message': f'User account permanently deleted from database'
            })
        else:
            return jsonify({'success': False, 'error': 'Database connection failed'}), 500
            
    except Exception as e:
        return jsonify({'success': False, 'error': f'Permanent deletion failed: {str(e)}'}), 500


@api_bp.route('/admin/permanently-delete-all-users', methods=['DELETE'])
@login_required
def permanently_delete_all_users():
    """Permanently delete all soft-deleted users from database (admin only)"""
    if not current_user.is_admin:
        return jsonify({'error': 'Admin access required'}), 403
    
    try:
        from firebase_models import User, Project, SurveyResponse
        from firebase_config import get_db
        
        # Get all deleted users
        deleted_users = User.get_deleted_users()
        
        if not deleted_users:
            return jsonify({'success': True, 'message': 'No deleted users to purge'})
        
        db = get_db()
        if not db:
            return jsonify({'success': False, 'error': 'Database connection failed'}), 500
        
        deleted_count = 0
        
        for user in deleted_users:
            # Don't allow deleting self
            if user.id == current_user.id:
                continue
                
            # Delete all user's projects
            projects = Project.get_by_user(user.id)
            for project in projects:
                project.delete()
            
            # Delete all user's survey responses
            survey_responses = SurveyResponse.get_by_user(user.id)
            for response in survey_responses:
                response.delete()
            
            # Permanently delete the user account from Firestore
            db.collection('users').document(user.id).delete()
            deleted_count += 1
        
        return jsonify({
            'success': True,
            'message': f'Successfully permanently deleted {deleted_count} user accounts from database'
        })
        
    except Exception as e:
        return jsonify({'success': False, 'error': f'Bulk permanent deletion failed: {str(e)}'}), 500


@api_bp.route('/export/<project_id>/<format>')
@login_required
def export_project(project_id, format):
    """Export project pitch deck in various formats"""
    if format == 'docx':
        return export_project_docx(project_id)
    elif format == 'pptx':
        return export_project_pptx(project_id)
    elif format == 'pdf':
        return export_project_pdf(project_id)
    else:
        return jsonify({'error': 'Unsupported format'}), 400

@api_bp.route('/generate-documentation/<project_id>/export', methods=['GET'])
@login_required
def export_documentation_docx(project_id):
    # Check if user has tokens for export
    if not current_user.can_export_files():
        return redirect(url_for('dashboard.out_of_tokens'))
    
    project = Project.get(project_id)
    if not project or project.user_id != current_user.id:
        return jsonify({'error': 'Project not found'}), 404
    
    # Create DOCX document
    doc = Document()
    
    # Title
    title = doc.add_heading(f"{project.title} - Technical Documentation", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Repository info
    doc.add_paragraph()
    doc.add_paragraph(f"Repository: {project.repo_owner}/{project.repo_name}")
    doc.add_paragraph(f"URL: {project.repo_url}")
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph()
    
    # Documentation sections
    if project.documentation:
        documentation = project.documentation
        
        sections = [
            ("Architecture Overview", documentation.get("architecture_overview")),
            ("API Documentation", documentation.get("api_documentation")),
            ("Database Schema", documentation.get("database_schema")),
            ("Deployment Guide", documentation.get("deployment_guide")),
            ("Configuration", documentation.get("configuration")),
            ("Testing", documentation.get("testing"))
        ]
        
        for section_title, content in sections:
            if content and content != "Not available in README.":
                doc.add_heading(section_title, level=1)
                doc.add_paragraph(content)
                doc.add_paragraph()
    else:
        doc.add_paragraph("Documentation not yet generated for this project.")
    
    # Save to buffer
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    # Deduct token for export
    current_user.use_token()
    
    filename = f"{project.title.replace(' ', '_')}_documentation.docx"
    return send_file(buffer, as_attachment=True, download_name=filename, 
                    mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

@api_bp.route('/generate-user-guide/<project_id>/export', methods=['GET'])
@login_required
def export_user_guide_docx(project_id):
    # Check if user has tokens for export
    if not current_user.can_export_files():
        return redirect(url_for('dashboard.out_of_tokens'))
    
    project = Project.get(project_id)
    if not project or project.user_id != current_user.id:
        return jsonify({'error': 'Project not found'}), 404
    
    # Create DOCX document
    doc = Document()
    
    # Title
    title = doc.add_heading(f"{project.title} - User Guide", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Repository info
    doc.add_paragraph()
    doc.add_paragraph(f"Repository: {project.repo_owner}/{project.repo_name}")
    doc.add_paragraph(f"URL: {project.repo_url}")
    doc.add_paragraph(f"Generated: {datetime.now().strftime('%B %d, %Y')}")
    doc.add_paragraph()
    
    # User guide sections
    if project.user_guide:
        user_guide = project.user_guide
        
        # Check if user guide has content field (new format) or individual sections (old format)
        if 'content' in user_guide:
            # New format - single content field
            content = user_guide.get('content', '')
            if content and content.strip():
                # Split content by markdown headers and add to document
                lines = content.split('\n')
                current_section = ""
                current_content = []
                
                for line in lines:
                    if line.startswith('**') and line.endswith('**'):
                        # Save previous section
                        if current_section and current_content:
                            doc.add_heading(current_section, level=1)
                            doc.add_paragraph('\n'.join(current_content))
                            doc.add_paragraph()
                        
                        # Start new section
                        current_section = line.strip('*').strip()
                        current_content = []
                    elif line.startswith('#'):
                        # Save previous section
                        if current_section and current_content:
                            doc.add_heading(current_section, level=1)
                            doc.add_paragraph('\n'.join(current_content))
                            doc.add_paragraph()
                        
                        # Start new section
                        current_section = line.lstrip('#').strip()
                        current_content = []
                    else:
                        if line.strip():
                            current_content.append(line)
                
                # Add final section
                if current_section and current_content:
                    doc.add_heading(current_section, level=1)
                    doc.add_paragraph('\n'.join(current_content))
                    doc.add_paragraph()
                
                # If no sections were found, add all content as one block
                if not current_section:
                    doc.add_heading("User Guide", level=1)
                    doc.add_paragraph(content)
                    doc.add_paragraph()
            else:
                doc.add_paragraph("User guide content is empty.")
        else:
            # Old format - individual sections
            sections = [
                ("Getting Started", user_guide.get("getting_started")),
                ("Installation", user_guide.get("installation")),
                ("Basic Usage", user_guide.get("basic_usage")),
                ("Advanced Features", user_guide.get("advanced_features")),
                ("Troubleshooting", user_guide.get("troubleshooting")),
                ("FAQ", user_guide.get("faq"))
            ]
            
            for section_title, content in sections:
                if content and content != "Not available in README.":
                    doc.add_heading(section_title, level=1)
                    doc.add_paragraph(content)
                    doc.add_paragraph()
    else:
        doc.add_paragraph("User guide not yet generated for this project.")
    
    # Save to buffer
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    # Deduct token for export
    current_user.use_token()
    
    filename = f"{project.title.replace(' ', '_')}_user_guide.docx"
    return send_file(buffer, as_attachment=True, download_name=filename, 
                    mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


@api_bp.route('/validate-coupon', methods=['POST'])
@login_required
def validate_coupon():
    """Validate coupon code and return discount information"""
    try:
        from firebase_models import Coupon
        from datetime import datetime
        
        data = request.get_json()
        coupon_code = data.get('code', '').upper().strip()
        plan_type = data.get('plan', 'basic')  # basic or pro
        
        if not coupon_code:
            return jsonify({'valid': False, 'error': 'Coupon code is required'}), 400
        
        # Get coupon by code
        coupon = Coupon.get_by_code(coupon_code)
        
        if not coupon:
            return jsonify({'valid': False, 'error': 'Invalid coupon code'}), 404
        
        # Check if coupon is valid
        if not coupon.is_valid():
            if not coupon.is_active:
                return jsonify({'valid': False, 'error': 'This coupon is no longer active'}), 400
            elif coupon.expires_at and datetime.utcnow() > coupon.expires_at:
                return jsonify({'valid': False, 'error': 'This coupon has expired'}), 400
            elif coupon.max_uses and coupon.current_uses >= coupon.max_uses:
                return jsonify({'valid': False, 'error': 'This coupon has reached its usage limit'}), 400
            else:
                return jsonify({'valid': False, 'error': 'This coupon is not valid'}), 400
        
        # Check if coupon applies to the selected plan
        if coupon.applies_to != 'all' and coupon.applies_to != plan_type:
            return jsonify({
                'valid': False, 
                'error': f'This coupon only applies to {coupon.applies_to} plans'
            }), 400
        
        # Calculate discount
        plan_prices = {'basic': 9.99, 'pro': 19.99}
        original_price = plan_prices.get(plan_type, 9.99)
        
        if coupon.discount_type == 'percentage':
            discount_amount = original_price * (coupon.discount_value / 100)
        else:
            discount_amount = min(coupon.discount_value, original_price)  # Don't exceed original price
        
        final_price = max(0, original_price - discount_amount)
        
        return jsonify({
            'valid': True,
            'coupon': {
                'code': coupon.code,
                'discount_type': coupon.discount_type,
                'discount_value': coupon.discount_value,
                'applies_to': coupon.applies_to
            },
            'pricing': {
                'original_price': original_price,
                'discount_amount': round(discount_amount, 2),
                'final_price': round(final_price, 2),
                'savings_percentage': round((discount_amount / original_price) * 100, 1)
            }
        })
    
    except Exception as e:
        return jsonify({'valid': False, 'error': f'Failed to validate coupon: {str(e)}'}), 500


@api_bp.route('/apply-coupon', methods=['POST'])
@login_required
def apply_coupon():
    """Apply coupon to user account for checkout"""
    try:
        from firebase_models import Coupon
        
        data = request.get_json()
        coupon_code = data.get('code', '').upper().strip()
        
        if not coupon_code:
            return jsonify({'success': False, 'error': 'Coupon code is required'}), 400
        
        # Validate coupon first
        coupon = Coupon.get_by_code(coupon_code)
        if not coupon or not coupon.is_valid():
            return jsonify({'success': False, 'error': 'Invalid or expired coupon'}), 400
        
        # Store coupon code in user's account for checkout
        current_user.coupon_code = coupon_code
        current_user.save()
        
        return jsonify({
            'success': True,
            'message': f'Coupon {coupon_code} applied successfully!',
            'coupon_code': coupon_code
        })
    
    except Exception as e:
        return jsonify({'success': False, 'error': f'Failed to apply coupon: {str(e)}'}), 500
